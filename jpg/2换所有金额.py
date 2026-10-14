# -*- coding: utf-8 -*-
"""JPG 完税证明金额替换脚本（极致对齐版）。

替换每张图片中的三处金额：
  1. 表格中"实缴(退)金额"列的数字（右对齐，带千分位逗号）
  2. "金额合计"行左侧的中文大写金额（楷体，左对齐）
  3. "金额合计"行右侧的 ¥ 数字：¥ 沿用原件自带的字形（不重画），只把其后的数字右对齐替换

对齐策略（全部以原图实测墨迹为准）：
  · 定位：先自动找表格横线/竖线，圈出三个单元格
  · 擦除：只擦掉原文字的墨迹矩形，不碰表格线（不会抹掉单元格竖线）
  · 字号：用同一字体，按原图该处墨迹高度反推字号
  · 位置：数字右边缘对齐原图右边缘；竖向上对齐原图墨迹中心

数据来源：jpg.xlsx（D列=姓名用于匹配文件名，H列=中文金额，I列=金额1，J列=金额2）
字体：大写金额用 STKAITI.TTF（楷体）；数字金额用 simhei.ttf（黑体）
      （SimHei 不含 ¥(U+00A5) 字形，会渲染成方框，故 ¥ 一律保留原件像素、不重画）
"""
import os
import sys

import numpy as np
import openpyxl
from PIL import Image, ImageDraw, ImageFont

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

SRC = os.path.dirname(os.path.abspath(__file__))
PARENT = os.path.dirname(SRC)
FONT_KAITI = os.path.join(SRC, 'STKAITI.TTF')                # 楷体：大写金额
FONT_HEI = r'C:\Windows\Fonts\simhei.ttf'                     # 黑体：数字金额
XLSX = os.path.join(SRC, 'jpg.xlsx')

LINE_TH = 200       # 表格线判定
INK_TH = 150        # 黑字判定：三通道都 < 150（红色印章因此被排除）
BG = (255, 255, 255)
FG = (0, 0, 0)
YEN_W = 13          # ¥ 字形墨迹宽度(px)：原件模板固定，¥ 直接沿用原件像素、不重画

# 字号：由原件实测反推（三张原件的模板一致，取固定值保证批量一致）
FONT_SIZE_NUM = 14  # SimHei 14px 时数字墨迹高 11，与原图"实缴(退)金额"数字吻合
FONT_SIZE_CH = 14   # STKAITI 14px 时，15/13/12 字墨迹宽 193/179/166，与原图 193/179/165 吻合

# 大写金额：原件已打印 "人民"，Excel 里是 "民币…"，补 "人" 后从第 3 个字"币"开始写
KEEP_PREFIX = 2


# ── 数据读取 ──────────────────────────────────────────
def load_xlsx(path):
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    data = {}
    for row in range(2, ws.max_row + 1):
        name = ws.cell(row, 4).value
        if not name:
            continue
        chinese = ws.cell(row, 8).value or ''
        amt1 = ws.cell(row, 9).value
        amt2 = ws.cell(row, 10).value
        data[str(name)] = {
            'chinese': str(chinese),
            'amt1': float(amt1) if amt1 is not None else 0.0,
            'amt2': float(amt2) if amt2 is not None else 0.0,
        }
    return data


def format_amount(val):
    s = f'{val:.2f}'
    i, d = s.split('.')
    return f'{int(i):,}.{d}'


# ── 表格线 / 单元格定位 ───────────────────────────────
def find_h_lines(gray, y1, y2):
    x1, x2 = 40, gray.shape[1] - 40
    need = (x2 - x1) * 0.9
    kept = []
    for y in range(y1, min(y2, gray.shape[0])):
        if np.sum(gray[y, x1:x2] < LINE_TH) > need:
            if not kept or y - kept[-1] > 2:
                kept.append(y)
    return kept


def find_v_lines(gray, y1, y2, frac=0.8):
    need = (y2 - y1) * frac
    return [x for x in range(2, gray.shape[1] - 2)
            if np.sum(gray[y1:y2, x] < LINE_TH) >= need]


def ink_box(rgb, x1, y1, x2, y2):
    """黑字墨迹矩形（含边框），返回 (left, top, right, bottom)"""
    reg = rgb[y1:y2, x1:x2]
    ys, xs = np.where(reg.max(axis=2) < INK_TH)
    if len(xs) == 0:
        return None
    return (x1 + int(xs.min()), y1 + int(ys.min()),
            x1 + int(xs.max()), y1 + int(ys.max()))


def ink_runs(rgb, x1, y1, x2, y2, gap=8):
    """按列把墨迹切成若干段（段间空白 > gap 才断开），返回每段的墨迹矩形"""
    m = rgb[y1:y2, x1:x2].max(axis=2) < INK_TH
    idx = np.where(m.any(axis=0))[0]
    if len(idx) == 0:
        return []
    runs, s, p = [], idx[0], idx[0]
    for c in idx[1:]:
        if c - p > gap:
            runs.append((s, p))
            s = c
        p = c
    runs.append((s, p))
    out = []
    for a, b in runs:
        ys = np.where(m[:, a:b + 1].any(axis=1))[0]
        out.append((x1 + int(a), y1 + int(ys.min()),
                    x1 + int(b), y1 + int(ys.max())))
    return out


def detect(rgb, gray):
    h = find_h_lines(gray, 280, 540)
    if len(h) < 3:
        raise ValueError(f'只找到 {len(h)} 条水平线，至少需要 3 条')
    table_top, table_bot = h[-3], h[-2]        # 表格数据行（"实缴(退)金额"所在行）
    total_top, total_bot = h[-2], h[-1]        # "金额合计"行

    ay1, ay2 = table_top + 2, table_bot - 1
    by1, by2 = total_top + 2, total_bot - 1

    data_vs = find_v_lines(gray, ay1, ay2)
    right = max(data_vs)                        # 右外边框
    border_left = min(x for x in data_vs if x > right - 4)     # 右边框最左列
    amt_left = max(x for x in data_vs if 60 < x < border_left - 3)   # 金额列左竖线

    total_vs = [x for x in find_v_lines(gray, by1, by2) if 60 < x < border_left - 3]
    label_left = min(total_vs)                  # "金额合计"标签右竖线
    yen_sep = max(total_vs)                     # 中文 / ¥ 的分隔竖线

    # 单元格可用范围：右端到 border_left-1。
    # border_left 是右边框竖线的抗锯齿列（灰），再往右才是线芯；原图数字墨迹
    # 其实会压到 border_left-1 这一列，若只截到 -2 会漏掉数字最后 1 列，
    # 且擦除擦不到它，留下一条灰残影，右端看起来"糊"。
    rx0, rx1 = amt_left + 2, border_left - 1    # 表格金额单元格 x 范围
    tx0, tx1 = yen_sep + 2, border_left - 1     # ¥ 单元格 x 范围
    cx0, cx1 = label_left + 2, yen_sep - 2      # 中文单元格 x 范围

    tb = ink_box(rgb, rx0, ay1, rx1 + 1, ay2 + 1)
    yn = ink_box(rgb, tx0, by1, tx1 + 1, by2 + 1)
    runs = ink_runs(rgb, cx0, by1, cx1 + 1, by2 + 1, gap=8)
    ch = max(runs, key=lambda b: b[2] - b[0]) if runs else None

    # ¥ 单元格里 ¥ 字形宽度固定(YEN_W)，右侧才是数字：只擦/换数字，¥ 保留原件像素
    yen_end = yn[0] + YEN_W - 1
    dn = ink_box(rgb, yen_end + 1, by1, tx1 + 1, by2 + 1)

    if tb is None or yn is None or ch is None or dn is None:
        raise ValueError(f'墨迹定位失败 tb={tb} yn={yn} ch={ch} dn={dn}')

    return {
        'tb': tb, 'yn': yn, 'ch': ch, 'dn': dn,
        'yen_end': yen_end,
        'tb_clip': (rx0, ay1, rx1, ay2),
        'dn_clip': (yen_end + 1, by1, tx1, by2),
        'ch_clip': (cx0, by1, cx1, by2),
        'h_lines': h,
    }


# ── 字体 / 渲染 ───────────────────────────────────────
_PAD = 24


def text_mask(font, text):
    """离屏渲染文字，返回 (贴图掩码, 画布内墨迹 bbox)。

    掩码用灰度作不透明度（保留抗锯齿）：贴回白底后与原渲染逐像素一致。
    墨迹范围由 <INK_TH 界定，与识别原图墨迹的阈值相同，定位一致。
    """
    w = int(font.getlength(text)) + _PAD * 2 + 8
    asc, desc = font.getmetrics()
    h = asc + desc + _PAD * 2 + 8
    canvas = Image.new('L', (max(w, 8), max(h, 8)), 255)
    ImageDraw.Draw(canvas).text((_PAD, _PAD), text, fill=0, font=font)
    g = np.array(canvas)
    ys, xs = np.where(g < INK_TH)
    bb = (int(xs.min()), int(ys.min()), int(xs.max()), int(ys.max()))
    return Image.fromarray((255 - g).astype(np.uint8)), bb


def draw_right(img, font, text, right_x, center_y):
    """右对齐（墨迹右边界 = right_x）；返回渲染后的墨迹左边界"""
    mask, (x0, y0, x1, y1) = text_mask(font, text)
    ax = right_x - x1
    ay = int(round(center_y - (y0 + y1) / 2.0))
    img.paste(FG, (int(round(ax)), ay), mask)
    return int(round(ax + x0))


def draw_left(img, font, text, left_x, center_y):
    """左对齐（墨迹左边界 = left_x）"""
    mask, (x0, y0, x1, y1) = text_mask(font, text)
    ax = left_x - x0
    ay = int(round(center_y - (y0 + y1) / 2.0))
    img.paste(FG, (int(round(ax)), ay), mask)


def erase(draw, box, clip):
    """白色擦除 box，并裁剪到 clip（不越出单元格 / 不碰表格线）"""
    x1, y1, x2, y2 = box
    cx0, cy0, cx1, cy1 = clip
    x1, x2 = max(x1, cx0), min(x2, cx1)
    y1, y2 = max(y1, cy0), min(y2, cy1)
    if x2 > x1 and y2 > y1:
        draw.rectangle([x1, y1, x2, y2], fill=BG)


# ── 主处理 ────────────────────────────────────────────
def process_image(img_path, info, out_path):
    img = Image.open(img_path).convert('RGB')
    rgb = np.array(img)
    gray = np.array(img.convert('L'))
    draw = ImageDraw.Draw(img)

    r = detect(rgb, gray)
    print(f'  水平线: {r["h_lines"]}')
    print(f'  表格数字: {r["tb"]}  中文: {r["ch"]}  ¥数字: {r["yn"]}')

    # 1) 表格数字金额（黑体）：右对齐到原图墨迹右边界，竖向对齐原图墨迹中心
    tb = r['tb']
    tb_text = format_amount(info['amt1'])
    tb_cy = (tb[1] + tb[3]) / 2.0
    erase(draw, (tb[0] - 1, tb[1] - 1, tb[2] + 1, tb[3] + 1), r['tb_clip'])
    f_tb = ImageFont.truetype(FONT_HEI, FONT_SIZE_NUM)
    draw_right(img, f_tb, tb_text, tb[2], tb_cy)
    print(f'  表格数字: "{tb_text}" 右边界={tb[2]} 原墨迹={tb}')

    # 2) 中文大写金额：原件已打印 "人民"，从第 KEEP_PREFIX 个字（"币"）开始写
    #    Excel 里存的是 "民币肆仟…"，补上开头 "人" 后为 "人民币肆仟…"
    ch = r['ch']
    full = '人' + info['chinese']
    keep = min(KEEP_PREFIX, len(full))
    f_ch = ImageFont.truetype(FONT_KAITI, FONT_SIZE_CH)
    # 在原图墨迹里按列切出各字形，取第 keep 个字（"币"）的起笔位置
    m = rgb[ch[1]:ch[3] + 1, ch[0]:ch[2] + 1].max(axis=2) < INK_TH
    groups, in_g, s = [], False, 0
    for i, d in enumerate(m.any(axis=0)):
        if d and not in_g:
            in_g, s = True, i
        elif not d and in_g:
            in_g = False
            groups.append((s, i - 1))
    if in_g:
        groups.append((s, m.shape[1] - 1))
    bi_start = ch[0] + groups[keep][0] if len(groups) > keep else ch[0]
    ch_text = full[keep:]
    ch_cy = (ch[1] + ch[3]) / 2.0
    erase(draw, (bi_start - 1, ch[1] - 1, ch[2] + 1, ch[3] + 1), r['ch_clip'])
    draw_left(img, f_ch, ch_text, bi_start, ch_cy)
    print(f'  中文大写: 保留"{full[:keep]}" 从"{ch_text}"开始 起笔={bi_start} 原墨迹={ch}')

    # 3) ¥ 数字（黑体）：¥ 沿用原件自带字形（不重画），只把 ¥ 之后的数字右对齐替换
    dn = r['dn']
    yn_text = format_amount(info['amt2'])
    dn_cy = (dn[1] + dn[3]) / 2.0
    erase(draw, (dn[0] - 1, dn[1] - 1, dn[2] + 1, dn[3] + 1), r['dn_clip'])
    f_yn = ImageFont.truetype(FONT_HEI, FONT_SIZE_NUM)
    num_left = draw_right(img, f_yn, yn_text, dn[2], dn_cy)
    print(f'  ¥数字:   "¥{yn_text}" ¥保留原件 右边界={dn[2]} 数字左边界={num_left} 原数字墨迹={dn}')

    # 校验：重新测量新图三处墨迹，与原图对比（中文窗口避开左侧"(大写)"与底线）
    out = np.array(img)
    vwin = {
        '表格数字': (r['tb_clip'][0], r['tb_clip'][1], r['tb_clip'][2], r['tb_clip'][3]),
        '中文大写': (max(r['ch_clip'][0], ch[0] - 3), r['ch_clip'][1] + 3,
                     r['ch_clip'][2], r['ch_clip'][3] - 3),
        '¥数字': (r['dn_clip'][0], r['dn_clip'][1], r['dn_clip'][2], r['dn_clip'][3]),
    }
    for label, box in (('表格数字', tb), ('中文大写', ch), ('¥数字', dn)):
        w = vwin[label]
        got = ink_box(out, w[0], w[1], w[2] + 1, w[3] + 1)
        if got:
            print(f'  校验 {label}: 原[{box[0]},{box[1]},{box[2]},{box[3]}] -> '
                  f'新[{got[0]},{got[1]},{got[2]},{got[3]}]')

    img.save(out_path, 'JPEG', quality=95)
    print(f'  已保存: {os.path.basename(out_path)}')
    return True


def main():
    data = load_xlsx(XLSX)
    print(f'读取 {len(data)} 条数据')

    jpgs = sorted(f for f in os.listdir(SRC)
                  if f.endswith('.jpg') and '_fixed' not in f)
    print(f'找到 {len(jpgs)} 张图片')

    success = 0
    for jpg in jpgs:
        parts = jpg.replace('.jpg', '').split('_')
        name = None
        for p in parts:
            p = p.strip()
            if '.' in p:
                p = p.split('.')[-1].strip()
            if len(p) >= 2 and not any(c.isdigit() for c in p) and p != '完税证明':
                name = p
                break

        if not name or name not in data:
            print(f'跳过 {jpg}：无法匹配人名 (提取到: {name})')
            continue

        info = data[name]
        print(f'\n处理: {jpg}')
        print(f'  姓名: {name}')

        in_path = os.path.join(SRC, jpg)
        out_path = os.path.join(SRC, jpg.replace('.jpg', '_fixed.jpg'))
        try:
            if process_image(in_path, info, out_path):
                success += 1
        except Exception as e:
            print(f'  错误: {e}')
            import traceback
            traceback.print_exc()

    print(f'\n完成: {success}/{len(jpgs)}')


if __name__ == '__main__':
    main()
