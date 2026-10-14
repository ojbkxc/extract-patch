# -*- coding: utf-8 -*-
"""
完税证明 PDF 原凭证号后三位重排 + PDF 改写工具
================================================================
遍历同目录下的 PDF，提取每份 PDF 中的"原凭证号"（18 位且以 332 开头，
排除 No● 后的票号），像 1读取pdf创建表格.py 换票号后三位那样，把原凭证号
的**后三位**用数字池随机重排（前 15 位不变）。

两件事一起做：
  1. 把换后的原凭证号写入 pdf.xlsx 的 L 列（纯文本，样式对齐 E 列）
  2. 就地修改 PDF 内容流中的原凭证号 TJ 数组（同 2换票号.py 原理），
     结果另存到 "已完成" 文件夹，原件不修改

换号算法与 1读取pdf创建表格.py 换票号完全同理论：
  * old = 原凭证号（18 位）
  * 数字池 pool = 二维码 parts[0]+parts[1]（区划码字段）里出现过的数字
  * pool 为空时回退用 old[:-3]（原凭证号前 15 位）的数字
  * rng = random.Random(seed|税号|old)，cands = pool 三位组合且 ≠ old[-3:]
  * 新原凭证号 = old[:-3] + rng.choice(cands)

用法：python 6换原凭证号.py [--seed N] [-v]
    --seed  随机种子，默认 20261013（同 1读取pdf创建表格.py，同种子+同输入可复现）
    -v      打印每行原凭证号变更明细
"""
import os
import re
import sys
import random
import argparse

import openpyxl
import pymupdf as fitz
from PIL import Image
import zxingcpp
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

BASE = os.path.dirname(os.path.abspath(__file__))
XLSX_NAME = "pdf.xlsx"
DONE_NAME = "已完成"
ORIG_COL = 12  # L 列

# 样式常量 —— 与 1读取pdf创建表格.py 完全一致
FONT_NAME = '宋体'
FILL_HEADER, FILL_ODD, FILL_EVEN = '1F4E79', 'FFFFFF', 'F7F9FC'
SIDE = Side(style='thin', color='D9DEE7')
BORDER = Border(left=SIDE, right=SIDE, top=SIDE, bottom=SIDE)
HEADER_H, DATA_H = 32.4, 28.8


def get_pdf_info(doc):
    """提取票号(No●后)、原凭证号(另一个332开头18位)、税号、二维码分段"""
    text = doc[0].get_text("text")
    m = re.search(r"No●(\d+)", text)
    cur_no = m.group(1) if m else None
    m2 = re.search(r"纳税人识别号\s*\n?\s*([0-9A-Za-z]+)", text)
    tax_no = m2.group(1).upper() if m2 else None
    # 原凭证号: 所有 332 开头的 18 位数字中，排除 No● 后的票号
    orig_no = None
    if cur_no:
        all_nums = re.findall(r"332\d{15}", text)
        candidates = [n for n in all_nums if n != cur_no]
        if candidates:
            orig_no = candidates[0]
    # 二维码解析 —— 取分段用于数字池（同 1读取pdf创建表格.py 途径）
    pix = doc[0].get_pixmap(dpi=300)
    img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    codes = zxingcpp.read_barcodes(img)
    qr = codes[0].text if codes else None
    parts = [p.strip() for p in qr.split(',')] if qr else []
    return cur_no, orig_no, tax_no, parts


def file_index(name):
    """文件名开头的数字，作为序号回退匹配用"""
    m = re.match(r'\s*(\d+)', name)
    return int(m.group(1)) if m else None


def style_header(cell):
    cell.font = Font(name=FONT_NAME, size=11, bold=True, color='FFFFFF')
    cell.fill = PatternFill('solid', fgColor=FILL_HEADER)
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = BORDER


def style_data(cell, row):
    i = row - 2
    cell.font = Font(name=FONT_NAME, size=11)
    cell.fill = PatternFill('solid', fgColor=FILL_ODD if i % 2 == 0 else FILL_EVEN)
    cell.border = BORDER
    cell.alignment = Alignment(horizontal='left', vertical='center')
    cell.number_format = '@'


def reshuffle_tail3(old, parts, tax_no, seed):
    """换后三位 —— 算法同 1读取pdf创建表格.py 换票号后三位。
    返回 (新号, 是否变更, 可用数字池, 是否回退)"""
    pool = sorted({ch for ch in (parts[0] + parts[1]) if ch.isdigit()}) if len(parts) >= 2 else []
    fallback = False
    if not pool and old and len(old) > 3:
        # 区划码字段无数字（如 null）：回退用原凭证号自身前 15 位的数字
        pool = sorted({ch for ch in old[:-3] if ch.isdigit()})
        fallback = bool(pool)
    if old and len(old) > 3 and pool:
        rng = random.Random('%s|%s|%s' % (seed, tax_no, old))
        cands = [a + b + c for a in pool for b in pool for c in pool if a + b + c != old[-3:]]
        return old[:-3] + rng.choice(cands), True, ''.join(pool), fallback
    return old, False, ''.join(pool), fallback


# ---- PDF 内容流就地替换（同 2换票号.py 原理）----
def _replace_in_stream(stream, old_b, new_b):
    matches = []
    for m in re.finditer(rb"\[([^\]]*)\]\s*TJ", stream):
        body = m.group(1)
        segs = [(mm.start(1), mm.end(1), mm.group(1))
                for mm in re.finditer(rb"\(([^)]*)\)", body)]
        if not segs:
            continue
        plain = b"".join(s for _, _, s in segs)
        if plain != old_b:
            continue
        new_body = bytearray(body)
        pos = 0
        repl = []
        for s, e, seg in segs:
            n = len(seg)
            repl.append((s, e, new_b[pos:pos + n]))
            pos += n
        for s, e, nb in reversed(repl):
            new_body[s:e] = nb
        matches.append((m.start(1), m.end(1), bytes(new_body)))
    if not matches:
        return None
    buf = bytearray(stream)
    for s, e, nb in reversed(matches):
        buf[s:e] = nb
    return bytes(buf)


def replace_in_pdf(doc, old_no, new_no):
    """就地替换 PDF 内容流中的原凭证号，返回是否替换成功"""
    if len(old_no) != len(new_no):
        raise ValueError("新旧号码长度不同，无法就地替换")
    old_b = old_no.encode("latin-1")
    new_b = new_no.encode("latin-1")
    changed = False
    for page in doc:
        for xref in page.get_contents():
            stream = doc.xref_stream(xref)
            new_stream = _replace_in_stream(stream, old_b, new_b)
            if new_stream is not None:
                doc.update_stream(xref, new_stream)
                changed = True
    return changed


def main():
    ap = argparse.ArgumentParser(description='原凭证号后三位重排 -> pdf.xlsx L 列 + 改写 PDF')
    ap.add_argument('--seed', type=int, default=20261013, help='随机种子，默认 20261013')
    ap.add_argument('-v', '--verbose', action='store_true', help='打印每行变更明细')
    a = ap.parse_args()

    xlsx_path = os.path.join(BASE, XLSX_NAME)
    if not os.path.exists(xlsx_path):
        print("找不到 %s，请把脚本放到 PDF 所在目录" % XLSX_NAME)
        return

    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active

    # 建立税号 -> 行号、序号 -> 行号 映射
    tax_to_row, idx_to_row = {}, {}
    for r in range(2, ws.max_row + 1):
        tax_no = ws.cell(r, 3).value
        if tax_no:
            tax_to_row[str(tax_no).strip().upper()] = r
        idx = ws.cell(r, 1).value
        if idx is not None:
            try:
                idx_to_row[int(idx)] = r
            except (TypeError, ValueError):
                pass

    # L 列表头（同 E 列表头写法）
    head = ws.cell(1, ORIG_COL)
    head.value = "原凭证号"
    style_header(head)
    ws.row_dimensions[1].height = HEADER_H

    done_dir = os.path.join(BASE, DONE_NAME)
    os.makedirs(done_dir, exist_ok=True)

    pdfs = sorted([f for f in os.listdir(BASE) if f.lower().endswith(".pdf")])
    ok = skip = fail = 0
    for name in pdfs:
        src = os.path.join(BASE, name)
        dst = os.path.join(done_dir, name)
        try:
            doc = fitz.open(src)
        except Exception as e:
            print("[失败] %s: 无法打开 %s" % (name, e))
            fail += 1
            continue

        cur_no, orig_no, tax_no, parts = get_pdf_info(doc)

        if not orig_no or not tax_no:
            print("[失败] %s: 未找到原凭证号或纳税人识别号" % name)
            fail += 1
            doc.close()
            continue

        # 匹配: 优先税号, 回退序号
        row = tax_to_row.get(str(tax_no).strip().upper())
        match_by = "税号"
        if not row:
            fi = file_index(name)
            if fi is not None:
                row = idx_to_row.get(fi)
                match_by = "序号"
        if not row:
            print("[失败] %s: xlsx 中无匹配(税号%s/序号%s)" % (name, tax_no, file_index(name)))
            fail += 1
            doc.close()
            continue

        # 换原凭证号后三位（同 1读取pdf创建表格.py 换票号后三位理论）
        new_orig, changed, pool, fallback = reshuffle_tail3(orig_no, parts, tax_no, a.seed)

        # 1) 写入 xlsx L 列（纯文本 + 样式）
        cell = ws.cell(row, ORIG_COL)
        style_data(cell, row)
        ws.row_dimensions[row].height = DATA_H
        cell.value = new_orig

        # 2) 改写 PDF 内容流，另存到 "已完成" 文件夹
        if not changed:
            doc.save(dst, garbage=4, deflate=True)
            doc.close()
            print("[跳过] %s: 无法确定可用数字，原凭证号未变更 %s" % (name, orig_no))
            skip += 1
            continue
        try:
            pdf_changed = replace_in_pdf(doc, orig_no, new_orig)
        except Exception as e:
            print("[失败] %s: PDF 改写异常 %s" % (name, e))
            fail += 1
            doc.close()
            continue
        if not pdf_changed:
            print("[失败] %s: PDF 内容流未找到原凭证号 %s" % (name, orig_no))
            fail += 1
            doc.close()
            continue
        doc.save(dst, garbage=4, deflate=True)
        doc.close()
        ok += 1
        if a.verbose:
            print("[完成] %s: 原凭证号后三位 %s -> %s  (可用数字 %s%s) | PDF 已改写"
                  % (name, orig_no[-3:], new_orig[-3:], pool,
                     '，区划码无数字回退用原凭证号前15位' if fallback else ''))
        else:
            print("[完成] %s: %s -> %s | PDF 已改写" % (name, orig_no, new_orig))

    wb.save(xlsx_path)
    print("\n总计: 完成 %d, 跳过 %d, 失败 %d, 共 %d 个 PDF" % (ok, skip, fail, len(pdfs)))
    print("已写回: %s (L列)" % xlsx_path)
    print("PDF 输出目录: %s" % done_dir)


if __name__ == "__main__":
    main()
