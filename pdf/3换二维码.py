"""
二维码替换脚本
读取 pdf.xlsx 的 E列(二维码的值)，按 C列(税号) 匹配 PDF，
用 qrcode 库生成新二维码，替换 PDF 中的原二维码图片。
输出到 "已完成" 文件夹。
"""
import os
import re
import glob
import io

import openpyxl
import qrcode
import fitz  # PyMuPDF

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
XLSX_PATH = os.path.join(BASE_DIR, "pdf.xlsx")
OUT_DIR = os.path.join(BASE_DIR, "已完成")
os.makedirs(OUT_DIR, exist_ok=True)


def load_xlsx_mapping():
    """读 Excel: C列(税号) -> E列(二维码的值)"""
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb.active
    mapping = {}
    for r in range(2, ws.max_row + 1):
        tax_id = ws.cell(r, 3).value
        qr_value = ws.cell(r, 5).value
        if tax_id and qr_value:
            mapping[str(tax_id).strip()] = str(qr_value).strip()
    return mapping


def generate_qr_png(data, scale=6):
    """用 qrcode 库生成二维码 PNG bytes，返回 bytes"""
    qr = qrcode.QRCode(
        version=1,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=scale,
        border=2,
    )
    qr.add_data(data)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def find_qr_xref(doc):
    """在 PDF 第1页找到二维码的 xref（最小的正方形 PNG 图片）"""
    page = doc[0]
    candidates = []
    for img in page.get_images(full=True):
        xref = img[0]
        pix = fitz.Pixmap(doc, xref)
        w, h = pix.width, pix.height
        pix = None
        if w == h and 30 <= w <= 500:
            stream = doc.xref_stream(xref)
            candidates.append((xref, w, h, len(stream)))
    if not candidates:
        return None
    candidates.sort(key=lambda c: c[3])
    return candidates[0][0]


def get_tax_id(doc):
    """从 PDF 文本中提取纳税人识别号"""
    text = doc[0].get_text()
    m = re.search(r"纳税人识别号[:\s]*([A-Z0-9]{15,20})", text)
    if m:
        return m.group(1).strip()
    page = doc[0]
    seq = []
    for block in page.get_text("dict")["blocks"]:
        if block["type"] != 0:
            continue
        for line in block["lines"]:
            for span in line["spans"]:
                seq.append(span)
    label = None
    for s in seq:
        if s["text"].strip() == "纳税人识别号":
            label = s
            break
    if not label:
        return None
    lx1 = label["bbox"][2]
    ly = label["origin"][1]
    for s in seq:
        if s["origin"][0] > lx1 and abs(s["origin"][1] - ly) < 20:
            m2 = re.search(r"([0-9A-Za-z]{15,20})", s["text"])
            if m2:
                return m2.group(1).upper()
    return None


def main():
    mapping = load_xlsx_mapping()
    print(f"Excel 中共 {len(mapping)} 条税号->二维码映射")

    pdfs = sorted(glob.glob(os.path.join(BASE_DIR, "*.pdf")))
    success = 0
    skip = 0

    for pdf_path in pdfs:
        fname = os.path.basename(pdf_path)
        doc = fitz.open(pdf_path)

        tax_id = get_tax_id(doc)
        if not tax_id or tax_id not in mapping:
            print(f"[跳过] {fname}: 税号 {tax_id} 未在 Excel 中找到")
            skip += 1
            doc.close()
            continue

        qr_value = mapping[tax_id]
        qr_xref = find_qr_xref(doc)
        if not qr_xref:
            print(f"[跳过] {fname}: 未找到二维码图片")
            skip += 1
            doc.close()
            continue

        png_bytes = generate_qr_png(qr_value, scale=6)
        page = doc[0]
        page.replace_image(qr_xref, stream=png_bytes)

        out_path = os.path.join(OUT_DIR, fname)
        doc.save(out_path, garbage=4, deflate=True)
        doc.close()
        print(f"[完成] {fname}: 税号={tax_id}, xref={qr_xref}")
        success += 1

    print(f"\n总计: {success} 成功, {skip} 跳过, {len(pdfs)} 个PDF")


if __name__ == "__main__":
    main()