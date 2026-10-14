# -*- coding: utf-8 -*-
"""
完税证明 PDF「计税依据」数字替换工具
================================================================
遍历同目录下的 PDF，提取每份 PDF 右下角"计税依据："后面的数字，
替换为 pdf.xlsx 表格 M 列（合计金额）的值，结果另存到"已完成"文件夹，
原件不修改。

计税依据数字的两种排版（换行逻辑）：
  * 同行：'备注: …  计税依据：180000'  —— 数字紧跟冒号后
  * 换行：'备注: …  计税依据：' \\n '210000'  —— 数字在下一行
  正则 r'计税依据[：:]\\s*([\\d.]+)' 的 \\s* 可匹配空格/换行，统一处理。

内容流替换原理（同 6换原凭证号.py，但支持新旧长度不同）：
  * 数字在 TJ 数组中被分割成多段，段间用 -1 位移隔开，如 (21)-1(00)-1(00)
  * 按旧段的字符长度分割新数字，保持段间位移不变，重构 TJ 数组 body
  * 旧数字 TJ 数组在每份 PDF 内唯一出现（经全量验证），通过值精确匹配

新值格式化：
  * 整数 -> 不带小数位（如 117623）
  * 非整数 -> 保留实际小数位，最多两位（如 117623.78、171629.7）

用法：python 7换计税依据.py [-v]
    -v  打印每行变更明细
"""
import os
import re
import sys
import argparse

import openpyxl
import pymupdf as fitz

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

BASE = os.path.dirname(os.path.abspath(__file__))
XLSX_NAME = "pdf.xlsx"
DONE_NAME = "已完成"
M_COL = 13  # M 列（合计金额）


def get_jiyi_num(text):
    """从 PDF 文本中提取"计税依据："后面的数字。

    处理同行（计税依据：180000）和换行（计税依据：\\n210000）两种排版。
    \\s* 可匹配空格、换行、制表符等，统一覆盖两种情况。
    返回数字字符串（如 '210000'、'147030.69'）或 None。
    """
    m = re.search(r'计税依据[：:]\s*([\d.]+)', text)
    return m.group(1) if m else None


def format_amount(val):
    """格式化 M 列金额值为字符串。

    整数 -> 不带小数位（如 117623）
    非整数 -> 保留实际小数位，最多两位（如 117623.78、171629.7）
    """
    if val is None:
        return None
    fval = round(float(val), 2)
    if fval == int(fval):
        return str(int(fval))
    s = '%.2f' % fval
    return s.rstrip('0').rstrip('.')


def file_index(name):
    """文件名开头的数字，作为序号回退匹配用"""
    m = re.match(r'\s*(\d+)', name)
    return int(m.group(1)) if m else None


# ---- PDF 内容流数字 TJ 数组替换 ----
def _replace_number_in_stream(stream, old_str, new_str):
    """在内容流中找到纯数字 TJ 数组（plain == old_str），替换为 new_str。

    TJ 数组中数字被分割成多段，段间有位移数字（如 -1）：
        旧 body: (21)-1(00)-1(00)   plain='210000'
    按旧段字符长度分割新数字，保持段间位移不变：
        新 body: (11)-1(76)-1(23.78)  plain='117623.78'

    返回替换后的新 stream（bytes），未找到则返回 None。
    """
    old_b = old_str.encode('latin-1')
    new_b = new_str.encode('latin-1')

    for m in re.finditer(rb'\[([^\]]*)\]\s*TJ', stream):
        body = m.group(1)
        # 提取所有 (字符串) 段的位置和内容
        segs = [(mm.start(1), mm.end(1), mm.group(1))
                for mm in re.finditer(rb'\(([^)]*)\)', body)]
        if not segs:
            continue
        plain = b''.join(s for _, _, s in segs)
        if plain != old_b:
            continue

        # 按旧段字符长度分割新数字：前 n-1 段保持旧长度，最后一段放剩余
        new_body = bytearray(body)
        pos = 0
        replacements = []
        for i, (s, e, old_seg) in enumerate(segs):
            n = len(old_seg) if i < len(segs) - 1 else len(new_b) - pos
            replacements.append((s, e, new_b[pos:pos + n]))
            pos += n
        for s, e, nb in reversed(replacements):
            new_body[s:e] = nb

        # 替换整个 TJ body
        buf = bytearray(stream)
        buf[m.start(1):m.end(1)] = bytes(new_body)
        return bytes(buf)
    return None


def replace_in_pdf(doc, old_num, new_num):
    """遍历所有页面内容流，替换计税依据数字 TJ 数组，返回是否替换成功"""
    changed = False
    for page in doc:
        for xref in page.get_contents():
            stream = doc.xref_stream(xref)
            new_stream = _replace_number_in_stream(stream, old_num, new_num)
            if new_stream is not None:
                doc.update_stream(xref, new_stream)
                changed = True
    return changed


def main():
    ap = argparse.ArgumentParser(description='替换 PDF 计税依据数字为 pdf.xlsx M 列值')
    ap.add_argument('-v', '--verbose', action='store_true', help='打印每行变更明细')
    a = ap.parse_args()

    xlsx_path = os.path.join(BASE, XLSX_NAME)
    if not os.path.exists(xlsx_path):
        print("找不到 %s，请先运行 1读取pdf创建表格.py 生成表格" % XLSX_NAME)
        return

    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active

    # 建立税号 -> 行号、序号 -> 行号 映射
    tax_to_row, idx_to_row = {}, {}
    for r in range(2, ws.max_row + 1):
        tax_no = ws.cell(r, 3).value
        if tax_no:
            tax_to_row[str(tax_no).strip().upper()] = r
        idx = ws.cell(r, 1).value
        if idx is not None:
            try:
                idx_to_row[int(idx)] = r
            except (TypeError, ValueError):
                pass

    done_dir = os.path.join(BASE, DONE_NAME)
    os.makedirs(done_dir, exist_ok=True)

    pdfs = sorted([f for f in os.listdir(BASE) if f.lower().endswith('.pdf')])
    ok = skip = fail = 0
    for name in pdfs:
        src = os.path.join(BASE, name)
        dst = os.path.join(done_dir, name)
        try:
            doc = fitz.open(src)
        except Exception as e:
            print("[失败] %s: 无法打开 %s" % (name, e))
            fail += 1
            continue

        text = doc[0].get_text('text')
        old_num = get_jiyi_num(text)

        if not old_num:
            print("[失败] %s: 未找到计税依据数字" % name)
            fail += 1
            doc.close()
            continue

        # 提取税号用于匹配 xlsx 行
        m = re.search(r'纳税人识别号\s*\n?\s*([0-9A-Za-z]+)', text)
        tax_no = m.group(1).upper() if m else None

        # 匹配: 优先税号, 回退序号
        row = tax_to_row.get(str(tax_no).strip().upper()) if tax_no else None
        if not row:
            fi = file_index(name)
            if fi is not None:
                row = idx_to_row.get(fi)
        if not row:
            print("[失败] %s: xlsx 中无匹配(税号%s/序号%s)" % (name, tax_no, file_index(name)))
            fail += 1
            doc.close()
            continue

        # 读取 M 列值并格式化
        m_val = ws.cell(row, M_COL).value
        new_num = format_amount(m_val)
        if not new_num:
            print("[失败] %s: M 列值为空" % name)
            fail += 1
            doc.close()
            continue

        # 替换 PDF 内容流中的计税依据数字
        try:
            pdf_changed = replace_in_pdf(doc, old_num, new_num)
        except Exception as e:
            print("[失败] %s: PDF 改写异常 %s" % (name, e))
            fail += 1
            doc.close()
            continue

        if not pdf_changed:
            print("[失败] %s: PDF 内容流未找到计税依据数字 %s" % (name, old_num))
            fail += 1
            doc.close()
            continue

        doc.save(dst, garbage=4, deflate=True)
        doc.close()
        ok += 1
        if a.verbose:
            print("[完成] %s: 计税依据 %s -> %s | PDF 已改写" % (name, old_num, new_num))
        else:
            print("[完成] %s: %s -> %s" % (name, old_num, new_num))

    print("\n总计: 完成 %d, 失败 %d, 共 %d 个 PDF" % (ok, fail, len(pdfs)))
    print("PDF 输出目录: %s" % done_dir)


if __name__ == '__main__':
    main()